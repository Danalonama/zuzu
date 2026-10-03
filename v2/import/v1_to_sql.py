#!/usr/bin/env python3
"""Turn the v1 Google Sheet (downloaded as .xlsx) into SQL for the v2 import.

    python3 v2/import/v1_to_sql.py zuzu-events.xlsx > v1_import_data.sql

Reads the tabs: events, venues, teachers, corrections. The output contains phone numbers
and other people's details — keep it out of git (v2/import/.gitignore covers it).

Mirrors how v1 shows events (apps-script/zuzu-events.gs, the "approved" feed):
  repeat=weekly    → weekday rule, end = date_end, or start + (count-1) weeks, or open-ended
  repeat=biweekly  → explicit dates every 14 days (v2 has no fortnightly rule)
  repeat=monthly   → explicit dates on the same nth weekday
  dates column     → explicit dates (wins over repeat, as in v1)
  teachers column  → "1.9=name; 8.9=name" → one child event per dated teacher
Two v1 quirks are deliberately NOT copied:
  - v1 reads a year-less "18.9" as the next 18.9 from *today*; here it is the one nearest
    the row's own start date (so a September course doesn't jump to next September).
  - v1 doesn't split dates on "·"; here "8.9 · 15.9" is two dates.

Requires: pip install openpyxl
"""
import datetime as dt
import json
import re
import sys

import openpyxl

TODAY = dt.date.today()
OPEN_HORIZON_DAYS = 90        # v1: open-ended repeats show ~90 days ahead
MONTHLY_HORIZON_DAYS = 120    # v1: open-ended monthly shows ~120 days ahead

AUDIENCE = {"נשים בלבד": "women_only", "גברים בלבד": "men_only", "לגברים בלבד": "men_only"}
SPELLING = {"חקר התנועה": "חקר תנועה"}
PLACEHOLDER_HOSTS = {"המורה יתעדכן"}
# "Moncalvo, Italy" in the city cell → an abroad venue, not an Israeli town
COUNTRIES = {"italy", "greece", "spain", "portugal", "france", "germany", "india", "netherlands", "uk", "usa"}


def rows(ws):
    it = ws.iter_rows(values_only=True)
    head = [str(h).strip() if h is not None else "" for h in next(it)]
    for r in it:
        if any(v not in (None, "") for v in r):
            yield dict(zip(head, r))


def s(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    t = s(v)
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d.%m.%Y"):
        try:
            return dt.datetime.strptime(t, fmt).date()
        except ValueError:
            pass
    return None


def as_time(v):
    """Pass times on as text; the database's parse_time() is the one parser."""
    if isinstance(v, dt.datetime):
        v = v.time()
    if isinstance(v, dt.time):
        return v.strftime("%H:%M")
    return s(v) or None


def day_month(p, anchor):
    """'18.9' / '18/9' / '18.9.26' / '2026-09-18' → date, year nearest the row's start."""
    p = p.strip()
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})$", p)
    if m:
        return dt.date(int(m[1]), int(m[2]), int(m[3]))
    m = re.match(r"^(\d{1,2})[./](\d{1,2})(?:[./](\d{2,4}))?$", p)
    if not m:
        return None
    day, mon = int(m[1]), int(m[2])
    try:
        if m[3]:
            yr = int(m[3])
            return dt.date(yr + 2000 if yr < 100 else yr, mon, day)
        base = anchor or TODAY
        cands = [dt.date(y, mon, day) for y in (base.year - 1, base.year, base.year + 1)]
        return min(cands, key=lambda d: abs((d - base).days))
    except ValueError:
        return None


def date_list(text, anchor):
    out = []
    for p in re.split(r"[,;\n·]+", text):
        d = day_month(p, anchor)
        if d:
            out.append(d)
    return sorted(set(out))


def nth_weekday(year, month, weekday, n):
    first = dt.date(year, month, 1)
    d = first + dt.timedelta(days=(weekday - first.weekday()) % 7 + 7 * (n - 1))
    return d if d.month == month else None


def monthly_dates(start, end, count):
    n = (start.day - 1) // 7 + 1
    out, y, m = [], start.year, start.month
    for _ in range(36):
        d = nth_weekday(y, m, start.weekday(), n)
        if d and d >= start:
            if count and len(out) >= count:
                break
            if not count and d > end:
                break
            out.append(d)
        m += 1
        if m > 12:
            m, y = 1, y + 1
    return out


def split(v):
    return [p.strip() for p in s(v).split(",") if p.strip()]


def venue_of(r):
    v = {"name": s(r.get("venue")) or None, "city": s(r.get("city")) or None,
         "region": s(r.get("region")) or None}
    tail = (v["city"] or "").rsplit(",", 1)
    if len(tail) == 2 and tail[1].strip().lower() in COUNTRIES:
        v.update(kind="abroad", country=tail[1].strip(), name=v["name"] or tail[0].strip(), city=None)
    return v


def convert(r):
    title = s(r.get("title"))
    start = as_date(r.get("date"))
    end = as_date(r.get("date_end"))
    repeat = s(r.get("repeat")).lower()
    count_txt = s(r.get("count"))
    count = int(count_txt) if count_txt.isdigit() else 0
    dates_txt = s(r.get("dates"))

    disciplines, formats, audience = [], [], []
    for v in split(r.get("category")):
        (audience.append(AUDIENCE[v]) if v in AUDIENCE else disciplines.append(SPELLING.get(v, v)))
    for v in split(r.get("type")):
        (audience.append(AUDIENCE[v]) if v in AUDIENCE else formats.append(v))

    hosts = [h for h in [s(r.get("host"))] if h and h not in PLACEHOLDER_HOSTS]
    price = s(r.get("price"))

    x = {
        "title": title,
        "date_start": start.isoformat() if start else None,
        "date_end": end.isoformat() if end else None,
        "time_start": as_time(r.get("time")),
        "time_end": as_time(r.get("time_end")),
        "hosts": hosts,
        "venue": venue_of(r),
        "disciplines": disciplines,
        "formats": formats,
        "audience": audience,
        "link": s(r.get("url")) or None,
        "phone": s(r.get("phone")) or None,
        "price_text": price or None,
    }
    if re.fullmatch(r"\d{1,6}", price):
        x["price_min"] = x["price_max"] = price

    open_ended = False
    if start and dates_txt:
        x["occurrence_dates"] = [d.isoformat() for d in date_list(dates_txt, start)]
    elif start and repeat == "weekly":
        x["weekdays"] = [(start.weekday() + 1) % 7]          # Python Monday=0 → Sunday=0
        if count:
            x["date_end"] = (start + dt.timedelta(weeks=count - 1)).isoformat()
    elif start and repeat == "biweekly":
        last = (start + dt.timedelta(weeks=2 * (count - 1)) if count
                else end or max(start, TODAY) + dt.timedelta(days=OPEN_HORIZON_DAYS))
        open_ended = not count and not end
        days, d = [], start
        while d <= last and len(days) < 200:
            days.append(d.isoformat())
            d += dt.timedelta(days=14)
        x["occurrence_dates"] = days
    elif start and repeat == "monthly":
        last = end or max(start, TODAY) + dt.timedelta(days=MONTHLY_HORIZON_DAYS)
        open_ended = not count and not end
        x["occurrence_dates"] = [d.isoformat() for d in monthly_dates(start, last, count)]

    rotation = []
    for part in re.split(r"[;\n]+", s(r.get("teachers"))):
        if "=" in part:
            d_txt, name = part.split("=", 1)
            d = day_month(d_txt, start)
            if d and name.strip() and name.strip() not in PLACEHOLDER_HOSTS:
                rotation.append({"date": d.isoformat(), "host": name.strip()})

    approved = r.get("approved") is True or s(r.get("approved")).upper() == "TRUE"
    rejected = s(r.get("review")).lower() == "rejected"
    lv = as_date(r.get("last_verified"))
    meta = {
        "uid": s(r.get("uid")) or None,
        "status": "live" if approved else "rejected" if rejected else "draft",
        "reason": s(r.get("reason")) or None,
        "open_ended": open_ended,
        "rotation": rotation,
        "row": {k: s(v) for k, v in r.items() if k and s(v)},
    }
    if lv:
        meta["last_verified"] = lv.isoformat()
    return meta, x


def lit(v):
    if v is None:
        return "null"
    if not isinstance(v, str):
        v = json.dumps(v, ensure_ascii=False)
    return "'" + v.replace("'", "''") + "'"


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = ["-- generated by v2/import/v1_to_sql.py from " + path.split("/")[-1],
           "-- contains people's names and phone numbers: do not commit, do not share",
           "begin;"]

    if "venues" in wb.sheetnames:
        for r in rows(wb["venues"]):
            if s(r.get("name")):
                v = venue_of({"venue": r.get("name"), "city": r.get("city"), "region": r.get("region")})
                region = "חו״ל" if v.get("kind") == "abroad" else v["region"]
                out.append(f"select v1_import_venue({lit(v['name'])}, {lit(v['city'])}, {lit(region)});")
    if "teachers" in wb.sheetnames:
        for r in rows(wb["teachers"]):
            if s(r.get("name")):
                out.append(f"select v1_import_teacher({lit(s(r.get('name')))}, {lit(s(r.get('phone')) or None)}, "
                           f"{lit(s(r.get('url')) or None)});")
    if "corrections" in wb.sheetnames:
        for r in rows(wb["corrections"]):
            if s(r.get("field")) in ("host", "venue"):
                out.append(f"select v1_learn_correction({lit(s(r.get('field')))}, {lit(s(r.get('from')))}, "
                           f"{lit(s(r.get('to')))});")

    events = list(rows(wb["events"]))
    events.sort(key=lambda r: s(r.get("timestamp")))
    seen = set()
    for i, r in enumerate(events):
        meta, x = convert(r)
        if not meta["uid"] or meta["uid"] in seen:      # a uid must be unique per source record
            meta["uid"] = f"row-{i + 2}"
        seen.add(meta["uid"])
        out.append(f"select v1_import_row({lit(meta)}::jsonb, {lit(x)}::jsonb);")

    out += ["select v1_import_finish();", "commit;", "select * from v1_import_report();"]
    sys.stdout.write("\n".join(out) + "\n")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
